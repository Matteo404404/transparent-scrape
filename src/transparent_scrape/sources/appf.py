"""APPF europarty contributions and donations (XLSX)."""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl

from transparent_scrape.core.http import download
from transparent_scrape.core.storage import save_json

# Known publication URLs (APPF cmsdata ids change per upload — override with --url)
PARTIES_XLSX_URLS: dict[int, str] = {
    2024: "https://www.appf.europa.eu/cmsdata/301299/2024%20PARTIES%20Contributions%20and%20Donations.xlsx",
    2023: "https://www.appf.europa.eu/cmsdata/302877/2023%20PARTIES%20Contributions%20and%20Donations_redacted.xlsx",
    2022: "https://www.appf.europa.eu/cmsdata/297603/European%20Political%20Parties%20Contributions%20and%20Donations%202022_Redacted.xlsx",
    2021: "https://www.appf.europa.eu/cmsdata/297619/PARTIES%20Contributions%20and%20donations%20related%20to%20financial%20year%202021_Redacted.xlsx",
    2020: "https://www.appf.europa.eu/cmsdata/297637/PARTIES%20Contributions%20and%20donations%20related%20to%20financial%20year%202020%20updated%202023-03-17_Redacted.xlsx",
}

DEFAULT_APPF_YEARS = [2020, 2021, 2022, 2023, 2024]

PARTIES_FILENAME = "appf_parties_{year}.xlsx"
PARTIES_JSON = "appf_parties_{year}.json"


def parties_xlsx_url(year: int, *, url: str | None = None) -> str:
    if url:
        return url
    if year not in PARTIES_XLSX_URLS:
        raise ValueError(f"no known APPF parties URL for {year}; pass --url")
    return PARTIES_XLSX_URLS[year]


def fetch_parties_xlsx(year: int, out_dir: Path, *, url: str | None = None) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    dest = out_dir / PARTIES_FILENAME.format(year=year)
    return download(parties_xlsx_url(year, url=url), dest)


def _num(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    return None


def parse_parties_summary(ws) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=8, values_only=True):
        party, abbr = row[0], row[1]
        if not party or not abbr:
            continue
        if isinstance(party, str) and party.startswith("European Political"):
            break
        rows.append(
            {
                "party": str(party).strip(),
                "abbreviation": str(abbr).strip(),
                "contributions_legal_persons_eur": _num(row[2]),
                "contributions_natural_persons_eur": _num(row[3]),
                "donations_legal_persons_eur": _num(row[4]),
                "donations_natural_persons_eur": _num(row[5]) if len(row) > 5 else None,
            }
        )
    return rows


def _parse_party_sheet(ws) -> dict[str, Any]:
    contributors: list[dict[str, Any]] = []
    donations: list[dict[str, Any]] = []
    section = None

    for row in ws.iter_rows(values_only=True):
        label = row[0] if row else None
        if not label:
            continue
        label_s = str(label).strip()

        if label_s == "Contributions":
            section = "contributions"
            continue
        if label_s == "Donations":
            section = "donations"
            continue
        if label_s.startswith("Sub-Total") or label_s.startswith("Total"):
            continue
        if label_s in ("Contributor", "Donor"):
            continue

        country = row[1] if len(row) > 1 else None
        amount = _num(row[2]) if len(row) > 2 else None

        entry = {
            "name": label_s,
            "country": str(country).strip() if country else None,
            "amount_eur": amount,
        }
        if section == "contributions" and label_s not in ("Contributions",):
            contributors.append(entry)
        elif section == "donations":
            donations.append(entry)

    return {"contributors": contributors, "donations": donations}


def _summary_sheet_name(wb: openpyxl.Workbook) -> str | None:
    for name in wb.sheetnames:
        if name.strip().lower() == "eupps all totals":
            return name
    return None


def _is_party_header(label: str) -> bool:
    raw = label.strip()
    if raw.startswith("Ø") or raw.startswith("\u00d8"):
        return True
    return False


def _clean_party_name(label: str) -> str:
    s = str(label).replace("\u00d8", "").replace("\u00a0", " ").strip()
    if s.startswith("Ø"):
        s = s[1:].strip()
    return s


def _parse_legacy_eupp_sheet(ws, year: int) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Single-sheet APPF layout used for 2020–2021."""
    summary: list[dict[str, Any]] = []
    party_details: dict[str, Any] = {}
    current_party: str | None = None
    section: str | None = None
    contributors: list[dict[str, Any]] = []
    donations: list[dict[str, Any]] = []
    contrib_legal: float | None = None
    contrib_natural: float | None = None
    don_legal: float | None = None
    don_natural: float | None = None

    def flush_party() -> None:
        nonlocal current_party, contributors, donations
        nonlocal contrib_legal, contrib_natural, don_legal, don_natural
        if not current_party:
            return
        key = current_party[:60]
        party_details[key] = {"contributors": contributors, "donations": donations}
        summary.append(
            {
                "party": current_party,
                "abbreviation": None,
                "contributions_legal_persons_eur": contrib_legal,
                "contributions_natural_persons_eur": contrib_natural,
                "donations_legal_persons_eur": don_legal,
                "donations_natural_persons_eur": don_natural,
            }
        )
        current_party = None
        contributors, donations = [], []
        contrib_legal = contrib_natural = don_legal = don_natural = None

    for row in ws.iter_rows(values_only=True):
        label = row[0] if row else None
        if label is None:
            continue
        label_s = str(label).strip()
        low = label_s.lower()

        if _is_party_header(label_s):
            flush_party()
            current_party = _clean_party_name(label_s)
            section = None
            continue
        if not current_party:
            continue
        if label_s == "Contributions":
            section = "contributions"
            continue
        if label_s in ("Donations", "Donations") or low == "donations":
            section = "donations"
            continue
        if low.startswith("sub-total contributions from legal"):
            contrib_legal = _num(row[2])
            continue
        if low == "individual contributions":
            contrib_natural = _num(row[2])
            continue
        if label_s in ("Contributor", "Donor", "Donor ", "-") or label_s.startswith("Contributor"):
            continue
        if label_s.startswith("Sub-") or low.startswith("the information"):
            continue
        if label_s in ("Total", "Total*") or (row[1] and str(row[1]).strip() in ("Total", "Total*")):
            amount = _num(row[2]) or _num(row[3])
            if section == "donations" and amount is not None:
                don_legal = amount
            elif section == "contributions" and amount is not None and contrib_legal is None:
                contrib_legal = amount
            continue

        country = row[2] if len(row) > 2 else None
        if section == "contributions" and label_s and country and str(country).strip() not in ("Country",):
            contributors.append(
                {
                    "name": label_s,
                    "country": str(country).strip() if country else None,
                    "amount_eur": None,
                }
            )
        elif section == "donations" and label_s and label_s not in ("Country", "Value"):
            amount = _num(row[3]) if len(row) > 3 else _num(row[2])
            donations.append(
                {
                    "name": label_s,
                    "country": str(country).strip() if country else None,
                    "amount_eur": amount,
                }
            )

    flush_party()
    return summary, party_details


def parse_parties_xlsx(path: Path, year: int) -> dict[str, Any]:
    wb = openpyxl.load_workbook(path, read_only=False, data_only=True)
    summary_name = _summary_sheet_name(wb)

    if summary_name:
        summary_ws = wb[summary_name]
        summary = parse_parties_summary(summary_ws)
        party_sheets: dict[str, Any] = {}
        skip = {summary_name}
        for name in wb.sheetnames:
            if name in skip:
                continue
            party_sheets[name] = _parse_party_sheet(wb[name])
        layout = "modern"
    elif "EUPP" in wb.sheetnames:
        summary, party_sheets = _parse_legacy_eupp_sheet(wb["EUPP"], year)
        layout = "legacy_eupp"
    else:
        raise ValueError(f"unknown APPF layout for {year}: {wb.sheetnames}")

    return {
        "source": "appf",
        "financial_year": year,
        "layout": layout,
        "fetched_at": datetime.now(timezone.utc).isoformat(),
        "xlsx_path": str(path),
        "summary": summary,
        "party_details": party_sheets,
    }


def fetch_and_parse_parties(year: int, out_dir: Path, *, url: str | None = None) -> Path:
    xlsx = fetch_parties_xlsx(year, out_dir, url=url)
    payload = parse_parties_xlsx(xlsx, year)
    json_path = out_dir / PARTIES_JSON.format(year=year)
    save_json(json_path, payload)
    return json_path


def fetch_and_parse_years(
    years: list[int],
    raw_dir: Path,
    parsed_dir: Path,
    *,
    urls: dict[int, str] | None = None,
) -> list[Path]:
    from transparent_scrape.core.storage import load_json as _load_json, update_manifest

    appf_dir = parsed_dir / "appf"
    appf_dir.mkdir(parents=True, exist_ok=True)
    paths: list[Path] = []
    urls = urls or {}
    for year in years:
        fetch_and_parse_parties(year, raw_dir, url=urls.get(year))
        data = _load_json(raw_dir / PARTIES_JSON.format(year=year))
        parsed_path = save_json(appf_dir / f"parties_{year}.json", data)
        paths.append(parsed_path)
    update_manifest(parsed_dir, "appf", {"years": years, "count": len(paths)})
    return paths
